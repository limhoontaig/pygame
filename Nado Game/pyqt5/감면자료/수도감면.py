import sys
import os
import re
import pandas as pd
from datetime import datetime
from PyQt5.QtWidgets import (
    QApplication, QMainWindow, QWidget, QLabel, QLineEdit, QPushButton,
    QGroupBox, QMessageBox, QFileDialog, QVBoxLayout, QHBoxLayout, QGridLayout, QListWidget
)
from PyQt5.QtCore import Qt

import warnings
warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')

class WaterCalcApp(QMainWindow):
    def __init__(self):
        super().__init__()
        self.init_paths()
        self.init_ui()

    def init_paths(self):
        now = datetime.now()
        yyyymm = now.strftime("%Y") + now.strftime("%m") + '월'
        yyyy = now.strftime("%Y")
        
        self.detail_file_path = f'D:/과장/1 1 부과자료/{yyyy}년/{yyyymm}/수도감면자료/2609_수도요금 감면세대 상세조회.xlsx'
        self.template_file_path = f'D:/과장/1 1 부과자료/{yyyy}년/Templates/Water_Template_File_for_XPERP_upload.xls'
        self.code_table_path = f'D:/과장/1 1 부과자료/{yyyy}년/Templates/xperp_code_comparasion_table.xlsx'
        self.save_dir_path = f'D:/과장/1 1 부과자료/{yyyy}년/{yyyymm}/xperp_감면자료'

    def init_ui(self):
        self.setWindowTitle('래미안개포루체하임아파트 세대 수도감면 관리 프로그램')
        self.resize(900, 720)  # 적정 창 크기 유지
        
        central_widget = QWidget(self)
        self.setCentralWidget(central_widget)

        main_layout = QVBoxLayout(central_widget)
        main_layout.setContentsMargins(15, 12, 15, 12)
        main_layout.setSpacing(8)

        # 타이틀 라벨
        self.label_title = QLabel('래미안개포루체하임아파트 세대 수도감면 관리 프로그램')
        font_title = self.label_title.font()
        font_title.setFamily('HY헤드라인M')
        font_title.setPointSize(14)
        self.label_title.setFont(font_title)
        self.label_title.setAlignment(Qt.AlignCenter)
        self.label_title.setFrameShape(QLabel.Panel)
        self.label_title.setLineWidth(3)
        main_layout.addWidget(self.label_title)

        def create_file_group(title, line_edit, btn_text, slot_func):
            gb = QGroupBox(title)
            layout = QHBoxLayout()
            layout.setContentsMargins(10, 5, 10, 5)
            line_edit.setReadOnly(False)
            layout.addWidget(line_edit)
            btn = QPushButton(btn_text)
            btn.setFixedWidth(110)
            btn.clicked.connect(slot_func)
            layout.addWidget(btn)
            gb.setLayout(layout)
            return gb

        self.le_file1 = QLineEdit(self.detail_file_path)
        main_layout.addWidget(create_file_group('수도요금 감면세대 상세조회 파일 선택', self.le_file1, '상세조회파일', self.select_detail_file))

        self.le_file2 = QLineEdit(self.template_file_path)
        main_layout.addWidget(create_file_group('XPERP Upload용 Template 파일 선택', self.le_file2, 'Template', self.select_template_file))

        self.le_file3 = QLineEdit(self.code_table_path)
        main_layout.addWidget(create_file_group('수도용 코드 조견표 파일', self.le_file3, 'Code Table', self.select_code_file))

        self.le_file4 = QLineEdit(self.save_dir_path)
        main_layout.addWidget(create_file_group('XPERP 할인자료 업로드파일 저장경로', self.le_file4, '저장경로', self.select_save_dir))

        def create_summary_group(title, prefix):
            gb = QGroupBox(title)
            grid = QGridLayout()
            grid.setContentsMargins(10, 4, 10, 4)
            grid.setSpacing(4)
            
            items = [
                ('다자녀감면', 0, 0), (f'{prefix}_dajanya', 0, 1),
                ('기초생활', 0, 2), (f'{prefix}_kicho', 0, 3),
                ('중증장애인', 0, 4), (f'{prefix}_jungjoong', 0, 5),
                ('유공자감면', 1, 0), (f'{prefix}_yugong', 1, 1),
                ('중복할인', 1, 2), (f'{prefix}_overlap', 1, 3),
                ('총 합 계', 1, 4), (f'{prefix}_total', 1, 5)
            ]
            
            inputs = {}
            for name, r, c in items:
                if name.startswith('le_'):
                    le = QLineEdit()
                    le.setReadOnly(True)
                    le.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
                    grid.addWidget(le, r, c)
                    inputs[name] = le
                else:
                    lbl = QLabel(name)
                    grid.addWidget(lbl, r, c)
            gb.setLayout(grid)
            return gb, inputs

        gb_bill, self.bill_inputs = create_summary_group('공제 종류별 총 공제요금 합계액 현황표 [고지서 기준]', 'le_bill')
        main_layout.addWidget(gb_bill)

        gb_act, self.act_inputs = create_summary_group('공제 종류별 총 공제요금 합계액 현황표 [실제 반영 기준]', 'le_act')
        main_layout.addWidget(gb_act)

        # 미반영(누락) 세대 내역 리스트박스 그룹 (높이 고정으로 여백 문제 해결)
        gb_unapplied = QGroupBox('템플릿 미반영(누락) 세대 내역 확인창')
        unapplied_layout = QVBoxLayout()
        unapplied_layout.setContentsMargins(10, 4, 10, 4)
        self.list_unapplied = QListWidget()
        self.list_unapplied.setFixedHeight(65)
        unapplied_layout.addWidget(self.list_unapplied)
        gb_unapplied.setLayout(unapplied_layout)
        main_layout.addWidget(gb_unapplied)

        # 하단 버튼 및 정보 레이아웃
        bottom_layout = QHBoxLayout()
        bottom_layout.setContentsMargins(0, 2, 0, 0)
        
        self.lbl_developer = QLabel('프로그램 작성 : 임훈택 Rev 2.4, 2026.10 Issued')
        bottom_layout.addWidget(self.lbl_developer)
        bottom_layout.addStretch()

        self.btn_verify = QPushButton('자료 검증')
        self.btn_verify.setFixedWidth(110)
        self.btn_verify.setFixedHeight(32)
        self.btn_verify.clicked.connect(lambda: QMessageBox.information(self, "안내", "계산 시작을 통해 검증이 반영됩니다."))
        bottom_layout.addWidget(self.btn_verify)

        self.btn_start = QPushButton('계산 시작')
        self.btn_start.setFixedWidth(130)
        self.btn_start.setFixedHeight(32)
        self.btn_start.setStyleSheet("font-weight: bold; color: blue;")
        self.btn_start.clicked.connect(self.start_calculation)
        bottom_layout.addWidget(self.btn_start)

        self.btn_close = QPushButton('종     료')
        self.btn_close.setFixedWidth(110)
        self.btn_close.setFixedHeight(32)
        self.btn_close.clicked.connect(self.close)
        bottom_layout.addWidget(self.btn_close)

        main_layout.addLayout(bottom_layout)

    def select_detail_file(self):
        fname, _ = QFileDialog.getOpenFileName(self, '상세조회 파일 선택', '', 'Excel Files (*.xlsx *.xls)')
        if fname: self.le_file1.setText(fname)

    def select_template_file(self):
        fname, _ = QFileDialog.getOpenFileName(self, 'Template 파일 선택', '', 'Excel Files (*.xls *.xlsx)')
        if fname: self.le_file2.setText(fname)

    def select_code_file(self):
        fname, _ = QFileDialog.getOpenFileName(self, '코드 조견표 파일 선택', '', 'Excel Files (*.xlsx *.xls)')
        if fname: self.le_file3.setText(fname)

    def select_save_dir(self):
        dname = QFileDialog.getExistingDirectory(self, '저장 경로 선택')
        if dname: self.le_file4.setText(dname)

    def load_water_code_mapping(self, code_file):
        df = pd.read_excel(code_file, sheet_name=0, header=None)
        mapping = {}
        for idx in range(20, min(25, len(df))):
            row = df.iloc[idx]
            val_cat = row.iloc[2]
            val_code = row.iloc[3]
            if pd.notna(val_cat) and pd.notna(val_code):
                mapping[str(val_cat).strip()] = str(val_code).strip()
        return mapping

    def get_mapped_code(self, cat_name, code_mapping):
        cat_str = str(cat_name)
        if '기초생활' in cat_str:
            return code_mapping.get('기초생활', '3')
        elif '다자녀' in cat_str:
            return code_mapping.get('다자녀', 'I')
        elif '독립유공' in cat_str or '유공자' in cat_str:
            return code_mapping.get('유공자', '2')
        elif '중증장애' in cat_str:
            return code_mapping.get('중증장애', 'T')
        return ''

    def start_calculation(self):
        f1 = self.le_file1.text()
        f2 = self.le_file2.text()
        f3 = self.le_file3.text()
        f4 = self.le_file4.text()

        if not os.path.exists(f1) or not os.path.exists(f2) or not os.path.exists(f3):
            QMessageBox.warning(self, "경고", "지정한 엑셀 파일 경로 중 존재하지 않는 파일이 있습니다.")
            return

        try:
            code_mapping = self.load_water_code_mapping(f3)

            bill_stats, df_detail_parsed = self.parse_detail_file(f1)
            
            self.bill_inputs['le_bill_dajanya'].setText(f"{bill_stats['다자녀감면']['cnt']}세대 / {int(bill_stats['다자녀감면']['amt']):>9,}")
            self.bill_inputs['le_bill_kicho'].setText(f"{bill_stats['기초생활']['cnt']}세대 / {int(bill_stats['기초생활']['amt']):>9,}")
            self.bill_inputs['le_bill_jungjoong'].setText(f"{bill_stats['중증장애']['cnt']}세대 / {int(bill_stats['중증장애']['amt']):>9,}")
            self.bill_inputs['le_bill_yugong'].setText(f"{bill_stats['유공자']['cnt']}세대 / {int(bill_stats['유공자']['amt']):>9,}")
            self.bill_inputs['le_bill_overlap'].setText(f"{bill_stats['중복']['cnt']}세대 / {int(bill_stats['중복']['amt']):>9,}")
            self.bill_inputs['le_bill_total'].setText(f"{bill_stats['총합']['cnt']}세대 / {int(bill_stats['총합']['amt']):>9,}")

            self.list_unapplied.clear()

            actual_stats, unapplied_records = self.apply_to_template(f2, f4, df_detail_parsed, code_mapping)
            
            self.act_inputs['le_act_dajanya'].setText(f"{actual_stats['다자녀감면']['cnt']}세대 / {int(actual_stats['다자녀감면']['amt']):>9,}")
            self.act_inputs['le_act_kicho'].setText(f"{actual_stats['기초생활']['cnt']}세대 / {int(actual_stats['기초생활']['amt']):>9,}")
            self.act_inputs['le_act_jungjoong'].setText(f"{actual_stats['중증장애']['cnt']}세대 / {int(actual_stats['중증장애']['amt']):>9,}")
            self.act_inputs['le_act_yugong'].setText(f"{actual_stats['유공자']['cnt']}세대 / {int(actual_stats['유공자']['amt']):>9,}")
            self.act_inputs['le_act_overlap'].setText(f"{actual_stats['중복']['cnt']}세대 / {int(actual_stats['중복']['amt']):>9,}")
            self.act_inputs['le_act_total'].setText(f"{actual_stats['총합']['cnt']}세대 / {int(actual_stats['총합']['amt']):>9,}")

            if unapplied_records:
                for rec in unapplied_records:
                    msg = f"[누락] 동: {rec['동']}호: {rec['호']} | 감면구분: {rec['감면구분']} | 금액: {rec['세대당금액']:,}원 (Template에 해당 동·호수 없음)"
                    self.list_unapplied.addItem(msg)
                QMessageBox.warning(self, "확인 필요", f"총 {len(unapplied_records)}세대가 Template 파일에 존재하지 않아 반영되지 않았습니다.\n하단 미반영 내역 창을 확인하세요.")
            else:
                self.list_unapplied.addItem("모든 감면 대상 세대가 템플릿에 정상적으로 반영되었습니다.")
                QMessageBox.information(self, "완료", "수도감면 계산 및 업로드 파일 생성이 성공적으로 완료되었습니다!")

        except Exception as e:
            import traceback
            err_msg = traceback.format_exc()
            print(f"[ERROR 상세]:\n{err_msg}")
            QMessageBox.critical(self, "에러 발생", f"처리 중 오류가 발생했습니다:\n{str(e)}")

    def parse_detail_file(self, detail_file):
        import openpyxl
        wb = openpyxl.load_workbook(detail_file, data_only=True)
        ws = wb.active

        bill_stats = {
            '다자녀감면': {'cnt': 0, 'amt': 0},
            '기초생활': {'cnt': 0, 'amt': 0},
            '중증장애': {'cnt': 0, 'amt': 0},
            '유공자': {'cnt': 0, 'amt': 0},
            '중복': {'cnt': 0, 'amt': 0},
            '총합': {'cnt': 0, 'amt': 0}
        }
        
        summary_units = {}
        for c in range(2, 7):
            cat = ws.cell(row=4, column=c).value
            amt_str = ws.cell(row=5, column=c).value
            cnt_str = ws.cell(row=6, column=c).value
            if cat:
                cat_clean = str(cat).replace('\n', ' ')
                amt_cleaned = re.sub(r'[^0-9]', '', str(amt_str)) if amt_str else '0'
                amt = int(amt_cleaned) if amt_cleaned else 0
                
                cnt = int(re.sub(r'[^0-9]', '', str(cnt_str))) if cnt_str else 0
                unit_amt = int(amt / cnt) if cnt > 0 else 0
                
                summary_units[cat_clean] = unit_amt
                
                if '다자녀' in cat_clean:
                    bill_stats['다자녀감면'] = {'cnt': cnt, 'amt': amt}
                elif '기초생활' in cat_clean:
                    bill_stats['기초생활'] = {'cnt': cnt, 'amt': amt}
                elif '중증장애' in cat_clean:
                    bill_stats['중증장애'] = {'cnt': cnt, 'amt': amt}
                elif '유공자' in cat_clean or '독립유공' in cat_clean:
                    bill_stats['유공자'] = {'cnt': cnt, 'amt': amt}

        total_cnt = sum([v['cnt'] for k, v in bill_stats.items() if k != '총합'])
        total_amt = sum([v['amt'] for k, v in bill_stats.items() if k != '총합'])
        bill_stats['총합'] = {'cnt': total_cnt, 'amt': total_amt}

        all_records = []
        current_row = 8
        while current_row <= ws.max_row:
            val = ws.cell(row=current_row, column=1).value
            if val and '□' in str(val) and '감면현황' in str(val):
                match_cat = re.search(r'□\s*(.*?)\s*감면현황', str(val))
                current_cat = match_cat.group(1).strip() if match_cat else ""
                
                data_row = current_row + 3
                while data_row <= ws.max_row:
                    r_val = ws.cell(row=data_row, column=1).value
                    if r_val is None or ('□' in str(r_val)):
                        break
                    
                    detail_text = ws.cell(row=data_row, column=3).value
                    if detail_text:
                        matches = re.findall(r'(\d+)[동\-]?\s*(\d+)호?', str(detail_text))
                        if matches:
                            dong_str, ho_str = matches[0]
                            unit_amt = summary_units.get(current_cat, 0)
                            all_records.append({
                                '동': int(dong_str),
                                '호': int(ho_str),
                                '감면구분': current_cat,
                                '세대당금액': unit_amt
                            })
                    data_row += 1
                current_row = data_row
            else:
                current_row += 1

        df_records = pd.DataFrame(all_records)
        return bill_stats, df_records

    def apply_to_template(self, template_file, save_dir, df_detail, code_mapping):
        df_template = pd.read_excel(template_file, skiprows=0)
        
        if '동' not in df_template.columns or '호' not in df_template.columns:
            raise ValueError("Template 파일에 '동' 또는 '호' 컬럼이 존재하지 않습니다.")

        df_template['동'] = pd.to_numeric(df_template['동'], errors='coerce').fillna(0).astype(int)
        df_template['호'] = pd.to_numeric(df_template['호'], errors='coerce').fillna(0).astype(int)
        
        unapplied_records = []

        if not df_detail.empty:
            df_detail['동'] = pd.to_numeric(df_detail['동'], errors='coerce').fillna(0).astype(int)
            df_detail['호'] = pd.to_numeric(df_detail['호'], errors='coerce').fillna(0).astype(int)

            grouped_records = []
            for (dong, ho), group in df_detail.groupby(['동', '호']):
                exists_in_template = not df_template[(df_template['동'] == dong) & (df_template['호'] == ho)].empty
                
                cats = group['감면구분'].tolist()
                unit_amounts = group['세대당금액'].tolist()
                
                if not exists_in_template:
                    for c_idx, cat_item in enumerate(cats):
                        unapplied_records.append({
                            '동': dong,
                            '호': ho,
                            '감면구분': cat_item,
                            '세대당금액': unit_amounts[c_idx]
                        })
                    continue

                codes = []
                calc_amounts = []
                for cat in cats:
                    code = self.get_mapped_code(cat, code_mapping)
                    if code: codes.append(code)
                    unit_amt = group.loc[group['감면구분'] == cat, '세대당금액'].values[0] if len(group.loc[group['감면구분'] == cat, '세대당금액'].values) > 0 else 0
                    calc_amounts.append(unit_amt)

                final_code = code_mapping.get('중복할인1', 'V') if len(codes) > 1 else (codes[0] if len(codes) == 1 else '')
                total_unit_amt = sum(calc_amounts)
                primary_cat = cats[0] if len(cats) == 1 else '중복'

                grouped_records.append({
                    '동': dong,
                    '호': ho,
                    '감면구분': final_code,
                    '수도감면금': int(total_unit_amt) if total_unit_amt != 0 else 0,
                    '원본카테고리': primary_cat,
                    '상세항목수': len(cats)
                })
            df_mapped = pd.DataFrame(grouped_records)
        else:
            df_mapped = pd.DataFrame(columns=['동', '호', '감면구분', '수도감면금', '원본카테고리', '상세항목수'])

        if '감면구분' in df_template.columns:
            df_template.drop(columns=['감면구분'], inplace=True, errors='ignore')
        if '수도감면금' in df_template.columns:
            df_template.drop(columns=['수도감면금'], inplace=True, errors='ignore')

        merged_df = pd.merge(df_template, df_mapped, how='left', on=['동', '호'])
        
        if '수도감면금' in merged_df.columns:
            merged_df['수도감면금'] = pd.to_numeric(merged_df['수도감면금'], errors='coerce').fillna(0).astype(int)
            
        if '감면구분' in merged_df.columns:
            merged_df['감면구분'] = merged_df['감면구분'].fillna('')

        applied_rows = merged_df[merged_df['수도감면금'] > 0]
        
        actual_stats = {
            '다자녀감면': {'cnt': 0, 'amt': 0},
            '기초생활': {'cnt': 0, 'amt': 0},
            '중증장애': {'cnt': 0, 'amt': 0},
            '유공자': {'cnt': 0, 'amt': 0},
            '중복': {'cnt': 0, 'amt': 0},
            '총합': {'cnt': 0, 'amt': 0}
        }
        
        for _, row in applied_rows.iterrows():
            amt = int(row['수도감면금'])
            cat = row.get('원본카테고리', '')
            cnt_items = row.get('상세항목수', 1)
            
            actual_stats['총합']['cnt'] += 1
            actual_stats['총합']['amt'] += amt
            
            if cnt_items > 1:
                actual_stats['중복']['cnt'] += 1
                actual_stats['중복']['amt'] += amt
            elif '다자녀' in str(cat):
                actual_stats['다자녀감면']['cnt'] += 1
                actual_stats['다자녀감면']['amt'] += amt
            elif '기초생활' in str(cat):
                actual_stats['기초생활']['cnt'] += 1
                actual_stats['기초생활']['amt'] += amt
            elif '중증장애' in str(cat):
                actual_stats['중증장애']['cnt'] += 1
                actual_stats['중증장애']['amt'] += amt
            elif '유공자' in str(cat) or '독립유공' in str(cat):
                actual_stats['유공자']['cnt'] += 1
                actual_stats['유공자']['amt'] += amt

        cols_to_drop = [c for c in ['원본카테고리', '상세항목수'] if c in merged_df.columns]
        if cols_to_drop:
            merged_df.drop(columns=cols_to_drop, inplace=True)

        for col in merged_df.columns:
            if col != '수도감면금':
                merged_df[col] = merged_df[col].fillna('')

        now = datetime.now()
        dt_str = now.strftime("%Y") + now.strftime("%m")
        file_name = dt_str + 'WATER_XPERP_Upload_i_columns.xlsx'
        
        if not os.path.exists(save_dir):
            os.makedirs(save_dir)
            
        save_path = os.path.join(save_dir, file_name)
        merged_df.to_excel(save_path, index=False, header=False)
        
        return actual_stats, unapplied_records

if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = WaterCalcApp()
    window.show()
    sys.exit(app.exec_())